import json
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from annotationxlsx.__main__ import main
from annotationxlsx.schema import COLUMNS, empty_row


def _ann(gene_id: str, **fields) -> dict:
    payload = {
        'gene_id': gene_id,
        'name': fields.get('name', gene_id),
        'function': fields.get('function'),
        'functional_category': fields.get('functional_category'),
        'drug_susc_impact': fields.get('drug_susc_impact'),
        'infection_impact': fields.get('infection_impact'),
        'go_terms': fields.get('go_terms'),
    }
    return payload


def test_json_file_directory_and_zip(tmp_path: Path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    one = tmp_path / 'one.json'
    one.write_text(json.dumps(_ann('Tb927.2.200', name='b', function='fb')))
    nested = tmp_path / 'dir' / 'nested'
    nested.mkdir(parents=True)
    (nested / 'a.json').write_text(json.dumps(_ann(
        'Tb927.1.100',
        name='a',
        function='fa',
        functional_category=['cat1', 'cat2'],
        go_terms=[{'id': 'GO:0005739', 'name': 'mitochondrion'}],
    )))
    zip_path = tmp_path / 'in.zip'
    with zipfile.ZipFile(zip_path, 'w') as handle:
        handle.writestr('output1/gen_c.json', json.dumps(_ann('Tb927.3.300', name='c')))

    workbook = tmp_path / 'cmp.xlsx'
    assert main(['json', '--input', str(one), '--workbook', str(workbook)]) == 0
    ws = load_workbook(workbook)['generated']
    assert [cell.value for cell in ws[1]] == list(COLUMNS)
    assert ws['A2'].value == 'Tb927.2.200'

    assert main(['json', '--input', str(tmp_path / 'dir'), '--workbook', str(workbook)]) == 0
    ws = load_workbook(workbook)['generated']
    assert ws['A2'].value == 'Tb927.1.100'
    assert ws['D2'].value == 'cat1; cat2'
    assert ws['G2'].value == 'GO:0005739 mitochondrion'

    assert main(['json', '--input', str(zip_path), '--workbook', str(workbook)]) == 0
    ws = load_workbook(workbook)['generated']
    assert ws['A2'].value == 'Tb927.3.300'


def test_json_rejects_xls(tmp_path: Path):
    src = tmp_path / 'a.json'
    src.write_text(json.dumps(_ann('Tb927.1.100')))
    assert main(['json', '--input', str(src), '--workbook', str(tmp_path / 'out.xls')]) == 1


def test_json_missing_input_is_hard_fail(tmp_path: Path):
    assert main(['json', '--input', str(tmp_path / 'missing.json'), '--workbook', str(tmp_path / 'out.xlsx')]) == 1


def test_veupath_appends_using_generated_order(tmp_path: Path, monkeypatch):
    from annotationxlsx import __main__ as cli
    from annotationxlsx.workbook import write_sheet

    workbook = tmp_path / 'cmp.xlsx'
    write_sheet(workbook, 'generated', [empty_row('B'), empty_row('A')])

    def fake_fetch(gene_ids, **kwargs):
        assert gene_ids == ['B', 'A']
        rows = []
        for gene_id in gene_ids:
            row = empty_row(gene_id)
            row['function'] = f'prod-{gene_id}'
            rows.append(row)
        return rows

    monkeypatch.setattr(cli, 'fetch_rows', fake_fetch)
    assert main(['veupath', '--workbook', str(workbook)]) == 0
    wb = load_workbook(workbook)
    assert wb.sheetnames == ['generated', 'veupathdb']
    assert wb['generated']['A2'].value == 'B'
    assert wb['veupathdb']['A2'].value == 'B'
    assert wb['veupathdb']['C2'].value == 'prod-B'
    assert wb['veupathdb']['A3'].value == 'A'
    assert wb['generated']['A2'].value == 'B'


def test_veupath_zip_fallback_sorted(tmp_path: Path, monkeypatch):
    from annotationxlsx import __main__ as cli

    zip_path = tmp_path / 'tbrucei_20261022.zip'
    with zipfile.ZipFile(zip_path, 'w') as handle:
        handle.writestr('gen_b.json', json.dumps(_ann('B')))
        handle.writestr('gen_a.json', json.dumps(_ann('A')))

    seen = {}

    def fake_fetch(gene_ids, **kwargs):
        seen['ids'] = gene_ids
        return [empty_row(gene_id) for gene_id in gene_ids]

    monkeypatch.setattr(cli, 'fetch_rows', fake_fetch)
    workbook = tmp_path / 'new.xlsx'
    assert main(['veupath', '--workbook', str(workbook), '--zip', str(zip_path)]) == 0
    assert seen['ids'] == ['A', 'B']
    assert load_workbook(workbook).sheetnames == ['veupathdb']


def test_veupath_http_failure_does_not_write(tmp_path: Path, monkeypatch):
    from annotationxlsx import __main__ as cli
    from annotationxlsx.errors import AnnotationXlsxError
    from annotationxlsx.workbook import write_sheet

    workbook = tmp_path / 'cmp.xlsx'
    write_sheet(workbook, 'generated', [empty_row('A')])
    before = workbook.read_bytes()

    def boom(gene_ids, **kwargs):
        raise AnnotationXlsxError('TriTrypDB request failed for every gene after retry')

    monkeypatch.setattr(cli, 'fetch_rows', boom)
    assert main(['veupath', '--workbook', str(workbook)]) == 1
    assert workbook.read_bytes() == before


def test_veupath_no_genes_hard_fail(tmp_path: Path):
    workbook = tmp_path / 'empty.xlsx'
    assert main(['veupath', '--workbook', str(workbook), '--zip', str(tmp_path / 'missing.zip')]) == 1
