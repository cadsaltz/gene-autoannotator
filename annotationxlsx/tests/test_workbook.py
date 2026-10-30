from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from annotationxlsx.errors import AnnotationXlsxError
from annotationxlsx.schema import COLUMNS, empty_row
from annotationxlsx.workbook import gene_ids_from_sheet, require_xlsx, write_sheet


def test_require_xlsx_rejects_xls(tmp_path: Path):
    with pytest.raises(AnnotationXlsxError):
        require_xlsx(tmp_path / 'out.xls')


def test_write_creates_workbook_with_header_and_row(tmp_path: Path):
    path = tmp_path / 'out.xlsx'
    row = empty_row('Tb927.1.100')
    row['name'] = 'foo'
    write_sheet(path, 'generated', [row])
    wb = load_workbook(path)
    assert wb.sheetnames == ['generated']
    ws = wb['generated']
    assert [cell.value for cell in ws[1]] == list(COLUMNS)
    assert ws['A2'].value == 'Tb927.1.100'
    assert ws['B2'].value == 'foo'
    assert ws.freeze_panes == 'A2'
    assert ws.auto_filter.ref is not None


def test_write_replaces_named_sheet_only(tmp_path: Path):
    path = tmp_path / 'out.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.title = 'keepme'
    ws['A1'] = 'stay'
    wb.create_sheet('generated')
    wb['generated']['A1'] = 'old'
    wb.save(path)
    write_sheet(path, 'generated', [empty_row('Tb927.2.200')])
    wb2 = load_workbook(path)
    assert 'keepme' in wb2.sheetnames
    assert wb2['keepme']['A1'].value == 'stay'
    assert wb2['generated']['A2'].value == 'Tb927.2.200'


def test_gene_ids_from_sheet_preserves_order(tmp_path: Path):
    path = tmp_path / 'out.xlsx'
    write_sheet(path, 'generated', [empty_row('B'), empty_row('A')])
    assert gene_ids_from_sheet(path, 'generated') == ['B', 'A']
    assert gene_ids_from_sheet(path, 'missing') is None


def test_gene_ids_from_sheet_rejects_nonexistent_xls(tmp_path: Path):
    with pytest.raises(AnnotationXlsxError):
        gene_ids_from_sheet(tmp_path / 'out.xls', 'generated')


def test_null_string_not_written(tmp_path: Path):
    path = tmp_path / 'out.xlsx'
    row = empty_row('Tb927.1.100')
    write_sheet(path, 'generated', [row])
    ws = load_workbook(path)['generated']
    assert ws['C2'].value in (None, '')


def test_unreadable_existing_file_raises(tmp_path: Path):
    path = tmp_path / 'out.xlsx'
    path.write_text('not a workbook')
    with pytest.raises(AnnotationXlsxError):
        write_sheet(path, 'generated', [empty_row('X')])
