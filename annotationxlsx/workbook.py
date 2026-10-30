from __future__ import annotations

import os
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from annotationxlsx.errors import AnnotationXlsxError
from annotationxlsx.schema import COLUMNS, truncate_cell

WRAP = Alignment(wrap_text=True, vertical='top')
HEADER_FONT = Font(bold=True)


def require_xlsx(path: Path) -> None:
    path = Path(path)
    if path.suffix.lower() != '.xlsx':
        raise AnnotationXlsxError(f'workbook must be a .xlsx file, got {path}')


def gene_ids_from_sheet(path: Path, sheet_name: str) -> list[str] | None:
    path = Path(path)
    if not path.is_file():
        return None
    require_xlsx(path)
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, KeyError, OSError) as exc:
        raise AnnotationXlsxError(f'not a readable .xlsx: {path}') from exc
    try:
        if sheet_name not in wb.sheetnames:
            return None
        ws = wb[sheet_name]
        rows = ws.iter_rows(values_only=True)
        try:
            header = [str(value).strip() if value is not None else '' for value in next(rows)]
        except StopIteration as exc:
            raise AnnotationXlsxError(f'sheet {sheet_name!r} has no header row') from exc
        if 'gene_id' not in header:
            raise AnnotationXlsxError(f'sheet {sheet_name!r} has no gene_id column')
        index = header.index('gene_id')
        ids = []
        for row in rows:
            if index >= len(row):
                continue
            value = row[index]
            if value is None or str(value).strip() == '':
                continue
            ids.append(str(value).strip())
        return ids
    finally:
        wb.close()


def write_sheet(
    path: Path,
    sheet_name: str,
    rows: list[dict[str, str]],
    *,
    warnings: list[str] | None = None,
) -> None:
    path = Path(path)
    require_xlsx(path)
    warnings = warnings if warnings is not None else []
    if path.exists():
        try:
            wb = load_workbook(path)
        except (InvalidFileException, BadZipFile, KeyError, OSError) as exc:
            raise AnnotationXlsxError(f'not a readable .xlsx: {path}') from exc
        if sheet_name in wb.sheetnames:
            del wb[sheet_name]
    else:
        wb = Workbook()
        default = wb.active
        wb.remove(default)
    ws = wb.create_sheet(sheet_name)
    for col, key in enumerate(COLUMNS, start=1):
        cell = ws.cell(1, col, key)
        cell.font = HEADER_FONT
        cell.alignment = WRAP
        ws.column_dimensions[get_column_letter(col)].width = 28
    for row_idx, row in enumerate(rows, start=2):
        for col, key in enumerate(COLUMNS, start=1):
            text, truncated = truncate_cell(row.get(key) or '')
            if truncated:
                warnings.append(f'truncated {key} for {row.get("gene_id")} at Excel cell limit')
            cell = ws.cell(row_idx, col, text if text else None)
            cell.alignment = WRAP
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    tmp = path.with_name(path.name + '.tmp')
    wb.save(tmp)
    os.replace(tmp, path)
