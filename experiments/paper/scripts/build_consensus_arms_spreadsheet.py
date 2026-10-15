#!/usr/bin/env python3
"""Build a per-arm team-review workbook for the consensus-arms pilot.

One workbook per arm: excerpt (needles highlighted from labels), the arm's trio
extractor outputs, the arm decision output, and column-G flags vs reviewer labels.

    PYTHONPATH=. python experiments/paper/scripts/build_consensus_arms_spreadsheet.py \
      --arm-dir experiments/paper/results/consensus-arms/<run>/arm_A \
      --labels experiments/paper/results/model_spread_15k/labels \
      --out ~/Downloads/consensus_arms_<run>_arm_A_team_review.xlsx

Excerpt/model text is never altered (only rich-text styling and flags are added); an
excerpt beyond Excel's 32,767-character cell limit is cut at the limit and noted in G.
"""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

REPO_ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(REPO_ROOT))

from experiments.paper.runners.consensus_arms_io import load_labels, packet_stem  # noqa: E402
from experiments.paper.scripts import consensus_arms_review as review  # noqa: E402

EXCEL_CELL_LIMIT = 32767

GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
RED_FILL = PatternFill("solid", fgColor="FFC7CE")
AMBER_FILL = PatternFill("solid", fgColor="FFEB9C")
GREY_FILL = PatternFill("solid", fgColor="E7E6E6")
NEEDLE_FILL = PatternFill("solid", fgColor="E2EFDA")
LABEL_FILL = PatternFill("solid", fgColor="D6E3F0")
TRIAL_FILL = PatternFill("solid", fgColor="2E75B6")
WRAP_TOP = Alignment(vertical="top", wrap_text=True)
THIN = Border(
    left=Side(style="thin", color="B0B0B0"), right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"), bottom=Side(style="thin", color="B0B0B0"),
)

LEGEND = (
    "HOW TO READ THIS SHEET (review overlay; excerpt and model text unchanged)\n"
    "Excerpt: GREEN BOLD DOUBLE-UNDERLINE = needle (fact an extractor should find).\n"
    "Outputs: green lines = real fact, orange = real but padded, red = false positive.\n"
    "Column G flags: ✔ real fact found, ◐ real but padded, ✖ false positive, "
    "○ missed needle, ✓ correctly empty.\n"
    "ARM OUTPUT values inherit the label of the trio candidate that proposed exactly that value; "
    "values no candidate proposed are marked UNLABELED."
)


def exact_json(obj: Any) -> str:
    return json.dumps(obj, ensure_ascii=False, indent=2)


def load_records(arm_dir: Path) -> list[dict[str, Any]]:
    path = arm_dir / "records.jsonl"
    if not path.is_file():
        raise SystemExit(f"missing records file: {path}")
    rows = [json.loads(line) for line in path.read_text().splitlines() if line.strip()]
    return sorted(rows, key=lambda r: r["trial_id"])


def load_manifest(arm_dir: Path) -> dict[str, Any]:
    path = arm_dir / "manifest.json"
    return json.loads(path.read_text()) if path.is_file() else {}


def arm_name(arm_dir: Path, manifest: dict[str, Any] | None = None) -> str:
    manifest = load_manifest(arm_dir) if manifest is None else manifest
    return str(manifest.get("arm") or arm_dir.name.removeprefix("arm_"))


def flag_fill(res: dict[str, Any]) -> PatternFill | None:
    classes = set(res["classes"].values())
    good = bool(classes & {"REAL", "PADDED"})
    bad = "FP" in classes
    if good and bad:
        return AMBER_FILL
    if bad:
        return RED_FILL
    if good:
        return GREEN_FILL
    if res["missed"]:
        return GREY_FILL
    return None


def _row_height(label: str, text: str) -> float:
    if label == "EXCERPT":
        return min(400, max(60, 12 + len(text) // 120))
    if "OUTPUT" in label or label.startswith("CANDIDATE"):
        return min(220, max(45, 12 + len(text) // 80))
    return 18


def write_block(
    ws, row: int, label: str, value: Any, *, flags: Any = None, flag_fill_: PatternFill | None = None,
    is_trial_header: bool = False,
) -> int:
    """Write label (A), merged value (B:F) and flags (G), mirroring build_team_review_spreadsheet."""
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
    cell_a = ws.cell(row, 1, label)
    cell_b = ws.cell(row, 2, "" if value is None else value)
    for col in range(1, 7):
        ws.cell(row, col).border = THIN
    cell_a.alignment = WRAP_TOP
    cell_b.alignment = WRAP_TOP
    if is_trial_header:
        for col in range(1, 7):
            ws.cell(row, col).fill = TRIAL_FILL
            ws.cell(row, col).font = Font(bold=True, color="FFFFFF")
        ws.row_dimensions[row].height = 22
    else:
        cell_a.fill = LABEL_FILL
        cell_a.font = Font(bold=True)
        ws.row_dimensions[row].height = _row_height(label, str(value))
    if flags is not None:
        cell_g = ws.cell(row, 7, flags)
        cell_g.alignment = WRAP_TOP
        if flag_fill_ is not None:
            cell_g.fill = flag_fill_
    return row + 1


def write_trial(ws, row: int, index: int, total: int, rec: dict[str, Any], arm: str,
                label_doc: dict[str, Any] | None) -> int:
    trial_id = rec["trial_id"]
    candidates: dict[str, Any] = rec.get("candidates") or {}
    row = write_block(ws, row, "TRIAL", f"{index} of {total} - {trial_id}", is_trial_header=True)
    row = write_block(ws, row, "trial_id", trial_id)

    # --- excerpt with needle highlights ---
    original = rec.get("excerpt_text") or ""
    text = original[:EXCEL_CELL_LIMIT]
    value: Any = text
    gflag: Any = None
    gfill = None
    if label_doc is None:
        gflag = "No label file for this trial (no needle highlights or flags)."
    else:
        needles = review.needles_of(label_doc)
        rt, located = review.styled_excerpt(text, [d["gold_quote"] for _, d in needles])
        if rt is not None:
            assert str(rt) == text  # characters untouched
            value = rt
        if needles:
            gflag = review.needle_context(text, needles, label_doc.get("excerpt_note", ""))
            gfill = NEEDLE_FILL
        else:
            gflag = f"NO NEEDLES in this excerpt.\nReviewer: {label_doc.get('excerpt_note', '')}"
    if len(original) > EXCEL_CELL_LIMIT:
        gflag = f"Excerpt is {len(original):,} chars; cut at Excel's {EXCEL_CELL_LIMIT:,}-char cell limit.\n" + (
            str(gflag) if gflag is not None else ""
        )
    row = write_block(ws, row, "EXCERPT", value, flags=gflag, flag_fill_=gfill)

    # --- trio candidates ---
    has_needles = bool(label_doc and review.needles_of(label_doc))
    for model, output in candidates.items():
        out_text = exact_json(output)
        rt_out: Any = out_text
        flags: Any = None
        fill = None
        if label_doc is not None:
            res = review.evaluate_candidate(label_doc, model, output)
            styled = review.styled_output(out_text, res["classes"])
            if styled is not None:
                assert str(styled) == out_text
                rt_out = styled
            flags = review.flag_text(res, has_needles) or None
            fill = flag_fill(res)
        row = write_block(ws, row, f"CANDIDATE {model}", rt_out, flags=flags, flag_fill_=fill)

    # --- arm decision output ---
    arm_output = rec.get("arm_output")
    out_text = exact_json(arm_output)
    rt_out = out_text
    flags = None
    fill = None
    if label_doc is not None:
        res = review.evaluate_arm_output(label_doc, candidates, arm_output)
        styled = review.styled_output(out_text, res["classes"])
        if styled is not None:
            assert str(styled) == out_text
            rt_out = styled
        flags = review.flag_text(res, has_needles) or None
        fill = flag_fill(res)
    row = write_block(ws, row, f"ARM OUTPUT ({arm})", rt_out, flags=flags, flag_fill_=fill)
    metrics = rec.get("decision_metrics")
    if metrics:
        row = write_block(ws, row, "decision_metrics", exact_json(metrics))
    return row + 1


def build_workbook(arm_dir: Path, labels_dir: Path, out_path: Path) -> Path:
    arm_dir, out_path = Path(arm_dir), Path(out_path)
    manifest = load_manifest(arm_dir)
    arm = arm_name(arm_dir, manifest)
    records = load_records(arm_dir)
    labels = load_labels(Path(labels_dir))

    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    n_labeled = sum(1 for r in records if packet_stem(r["trial_id"]) in labels)
    readme_rows = [
        ("Workbook", str(out_path)),
        ("Arm", arm),
        ("Arm dir", str(arm_dir)),
        ("Trio models", ", ".join(manifest.get("models") or [])),
        ("Judge model", manifest.get("judge_model")),
        ("Classifier", manifest.get("classifier")),
        ("Run status", manifest.get("status")),
        ("Dry run", manifest.get("dry_run")),
        ("Trials", f"{len(records)} ({n_labeled} with reviewer labels)"),
        ("Legend", LEGEND),
    ]
    for r, (key, val) in enumerate(readme_rows, start=1):
        readme.cell(r, 1, key).font = Font(bold=True)
        cell = readme.cell(r, 2, "" if val is None else str(val))
        cell.alignment = WRAP_TOP
    readme.column_dimensions["A"].width = 18
    readme.column_dimensions["B"].width = 110

    ws = wb.create_sheet(f"Arm {arm}"[:31])
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 140
    for col in "CDEF":
        ws.column_dimensions[col].width = 12
    ws.column_dimensions["G"].width = 70
    row = 1
    for index, rec in enumerate(records, start=1):
        row = write_trial(ws, row, index, len(records), rec, arm, labels.get(packet_stem(rec["trial_id"])))
    ws.freeze_panes = "A2"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def build_parser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    p.add_argument("--arm-dir", required=True, type=Path, help="<run>/arm_<name> with records.jsonl")
    p.add_argument("--labels", required=True, type=Path, help="directory of t*.json reviewer labels")
    p.add_argument("--out", required=True, type=Path, help="output .xlsx path")
    return p


def main(argv: list[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    out = build_workbook(args.arm_dir.expanduser(), args.labels.expanduser(), args.out.expanduser())
    print(f"wrote {out} ({out.stat().st_size} bytes)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
