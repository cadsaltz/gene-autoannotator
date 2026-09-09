#!/usr/bin/env python3
"""Build a team-review Excel workbook from one tie-break consensus run directory."""

from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

REPO_ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(REPO_ROOT))

from autoannotation import llms  # noqa: E402

# Same text as experiments.paper.runners.general_consensus (kept inline so this
# script works without importing the experiments-branch runner).
GENERAL_BATCH_CONSENSUS_PROMPT = """
You merge candidate answers from different extractor models.

Candidate objects:
{candidates_json}

Merge ONLY these fields: {field_list}

Return JSON with exactly those keys. Use null for any field you cannot
reconcile from the candidates.

Rules:
- Reconcile only from the candidate values provided. You do not have access to
  the source text.
- When candidates agree exactly or as paraphrases, return concise wording drawn
  from those candidates.
- Prefer a value supported by a clear majority.
- Do not invent or add facts absent from every candidate.
- Return null for irreconcilable conflicts with no clear majority.
"""

NON_NONSENSE_TAG = "tiebreak-non-nonsense"
NONSENSE_TAG = "tiebreak-nonsense"
SHEET_BY_TAG = (
    (NON_NONSENSE_TAG, "Non-nonsense"),
    (NONSENSE_TAG, "Nonsense"),
)
# Mirror run_tiebreak_consensus.GENERAL_MULTI_SLOTS (kept inline to avoid importing
# the runner from this optional openpyxl script).
GENERAL_MULTI_SLOTS = ("function", "drug_susc_impact", "infection_impact")

header_font = Font(bold=True, color="FFFFFF")
label_fill = PatternFill("solid", fgColor="D6E3F0")
label_font = Font(bold=True)
trial_fill = PatternFill("solid", fgColor="2E75B6")
thin = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)


def exact_json(obj) -> str:
    return json.dumps(obj, ensure_ascii=False, indent=2)


def load_jsonl(path: Path) -> list[dict]:
    return [
        json.loads(line)
        for line in path.read_text().splitlines()
        if line.strip()
    ]


def load_aggregate(path: Path) -> list[dict]:
    if not path.is_file():
        return []
    with path.open(newline="") as handle:
        return list(csv.DictReader(handle))


def pair_cases(records: list[dict]) -> list[dict]:
    by_case: dict[str, dict] = {}
    for row in records:
        case = by_case.setdefault(row["case_id"], {"case_id": row["case_id"]})
        case[row["condition"]] = row
    ordered = []
    for case_id, bundle in by_case.items():
        hybrid = bundle.get("hybrid_consensus")
        if hybrid is None:
            raise ValueError(f"missing hybrid_consensus for {case_id}")
        ordered.append(
            {
                "hybrid": hybrid,
                "baseline": bundle.get("no_consensus_pick_extractor_0"),
            }
        )
    ordered.sort(key=lambda item: item["hybrid"]["case_id"])
    return ordered


def infer_experiment_tags(record: dict) -> list[str]:
    """Return stored tags, or infer for older records that lack experiment_tags."""
    tags = list(record.get("experiment_tags") or [])
    if tags:
        return tags
    case_family = str(record.get("case_family") or "")
    case_id = str(record.get("case_id") or "")
    if "nonsense" in case_family or "nonsense" in case_id:
        return [NONSENSE_TAG]
    return [NON_NONSENSE_TAG]


def filter_records_by_tag(records: list[dict], tag: str) -> list[dict]:
    return [row for row in records if tag in infer_experiment_tags(row)]


def field_keys_for_prompt(hybrid: dict) -> list[str]:
    raw = hybrid.get("field_keys")
    if isinstance(raw, list) and raw:
        return [str(key) for key in raw]
    field_key = str(hybrid.get("field_key") or "")
    if "," in field_key:
        return [part.strip() for part in field_key.split(",") if part.strip()]
    return [field_key] if field_key else []


def consensus_prompt_for_case(hybrid: dict) -> str | None:
    """Rebuild the prompt the live merger would send when LLM is invoked.

    General single-field cases are remapped onto the temporary ``function``
    slot; general multi-field cases map onto GENERAL_MULTI_SLOTS (see
    run_tiebreak_consensus._multi_field_case_records).
    """
    if not hybrid.get("llm_invoked"):
        return None

    field_keys = field_keys_for_prompt(hybrid)
    if not field_keys:
        return None
    candidates = hybrid["candidates"]
    domain = hybrid["domain"]

    if domain == "general":
        if len(field_keys) == 1:
            payload = [
                {"function": candidate.get(field_keys[0])} for candidate in candidates
            ]
            field_list = "function"
        else:
            if len(field_keys) > len(GENERAL_MULTI_SLOTS):
                raise ValueError(
                    f"{hybrid.get('case_id')}: general multi-field supports at most "
                    f"{len(GENERAL_MULTI_SLOTS)} fields"
                )
            slot_by_field = {
                field_key: GENERAL_MULTI_SLOTS[index]
                for index, field_key in enumerate(field_keys)
            }
            payload = [
                {
                    slot_by_field[field_key]: candidate.get(field_key)
                    for field_key in field_keys
                }
                for candidate in candidates
            ]
            field_list = ",".join(slot_by_field[key] for key in field_keys)
        return GENERAL_BATCH_CONSENSUS_PROMPT.format(
            candidates_json=json.dumps(payload, indent=2, ensure_ascii=False),
            field_list=field_list,
        ).strip()

    payload = [
        {field_key: candidate.get(field_key) for field_key in field_keys}
        for candidate in candidates
    ]
    return llms.BATCH_CONSENSUS_PROMPT.format(
        candidates_json=json.dumps(payload, indent=2, ensure_ascii=False),
        field_list=",".join(field_keys),
    ).strip()


def path_flag(hybrid: dict) -> str:
    if hybrid.get("llm_invoked"):
        return f"llm ({hybrid.get('provenance', 'unknown')})"
    return f"deterministic ({hybrid.get('provenance', 'unknown')})"


def row_height(label: str, value) -> float:
    text = "" if value is None else str(value)
    if "PROMPT" in label:
        return min(400, max(60, 12 + len(text) // 120))
    if "CANDIDATES" in label or "OUTPUT" in label or "RESULT" in label:
        return min(220, max(45, 12 + len(text) // 80))
    return 18


def write_block(ws, row: int, label: str, value, *, is_trial_header: bool = False) -> int:
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
    cell_a = ws.cell(row, 1, label)
    cell_b = ws.cell(row, 2, "" if value is None else str(value))
    cell_a.border = thin
    for col in range(2, 7):
        ws.cell(row, col).border = thin
    cell_a.alignment = Alignment(vertical="top", wrap_text=True)
    cell_b.alignment = Alignment(vertical="top", wrap_text=True)
    if is_trial_header:
        cell_a.fill = trial_fill
        cell_a.font = header_font
        for col in range(2, 7):
            cell = ws.cell(row, col)
            cell.fill = trial_fill
            cell.font = header_font
        ws.row_dimensions[row].height = 22
    else:
        cell_a.fill = label_fill
        cell_a.font = label_font
        ws.row_dimensions[row].height = row_height(label, value)
    return row + 1


def write_cases(ws, cases: list[dict], *, consensus_model: str) -> None:
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 140
    for col in ("C", "D", "E", "F"):
        ws.column_dimensions[col].width = 12

    row = 1
    for index, bundle in enumerate(cases, start=1):
        hybrid = bundle["hybrid"]
        baseline = bundle.get("baseline") or {}
        prompt = consensus_prompt_for_case(hybrid)
        if prompt is None:
            prompt_text = "N/A — deterministic majority (no LLM call)"
        else:
            prompt_text = prompt

        blocks: list[tuple[str, object, bool]] = [
            (
                "CASE",
                f'{index} of {len(cases)} — {hybrid["case_id"]}',
                True,
            ),
            ("case_id", hybrid["case_id"], False),
            ("domain", hybrid.get("domain"), False),
            ("case_family", hybrid.get("case_family"), False),
            ("experiment_tags", exact_json(hybrid.get("experiment_tags") or []), False),
            ("field_key", hybrid.get("field_key"), False),
            ("expected", exact_json(hybrid.get("expected")), False),
            ("CANDIDATES (extractor 0/1/2)", exact_json(hybrid.get("candidates")), False),
            ("CONSENSUS PROMPT (what consensus sees)", prompt_text, False),
            (
                f"CONSENSUS RESULT ({consensus_model})",
                exact_json(hybrid.get("observed")),
                False,
            ),
            ("consensus_path", path_flag(hybrid), False),
            ("provenance", hybrid.get("provenance"), False),
            ("llm_invoked", hybrid.get("llm_invoked"), False),
            ("match_exact", hybrid.get("match_exact"), False),
            ("match_soft", hybrid.get("match_soft"), False),
            ("invention", hybrid.get("invention"), False),
            (
                "baseline (extractor 0) observed",
                exact_json(baseline.get("observed")),
                False,
            ),
            (
                "baseline match_soft / match_exact",
                f"{baseline.get('match_soft')} / {baseline.get('match_exact')}",
                False,
            ),
        ]
        for label, value, is_header in blocks:
            row = write_block(ws, row, label, value, is_trial_header=is_header)
        row += 1

    ws.freeze_panes = "A2"


METRIC_KEYS = (
    ("case_count", "case_count"),
    ("consensus_match_soft", "consensus_match_soft_rate"),
    ("consensus_match_exact", "consensus_match_exact_rate"),
    ("baseline_match_soft", "baseline_match_soft_rate"),
    ("necessity_delta", "necessity_delta"),
    ("llm_invoked_rate", "llm_invoked_rate"),
    ("expect_llm_calibration_rate", "expect_llm_calibration_rate"),
    ("invention_rate", "invention_rate"),
    ("nonsense_majority_adoption_rate", "nonsense_majority_adoption_rate"),
)


def overall_aggregate(rows: list[dict]) -> dict:
    for row in rows:
        if row.get("scope") == "overall" and row.get("value") == "all":
            return row
    return rows[0] if rows else {}


def tag_aggregate(rows: list[dict], tag: str) -> dict:
    for row in rows:
        if row.get("scope") == "experiment_tag" and row.get("value") == tag:
            return row
    return {}


def metric_summary_rows(prefix: str, aggregate: dict) -> list[tuple[str, object]]:
    return [
        (f"{prefix}.{label}", aggregate.get(key))
        for label, key in METRIC_KEYS
    ]


def write_summary(
    ws,
    *,
    manifest: dict,
    run_dir: Path,
    aggregate: list[dict],
    present_tags: list[str],
) -> None:
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 80
    ws.cell(1, 1, "Tie-break consensus — team review")
    ws.cell(1, 1).font = Font(bold=True, size=14)
    ws.merge_cells("A1:B1")

    consensus_model = (manifest.get("model_tags") or {}).get("consensus", "?")
    overall = overall_aggregate(aggregate)

    sections: list[tuple[str, list[tuple[str, object]]]] = [
        (
            "Run",
            [
                ("experiment_id", manifest.get("experiment_id")),
                ("run_id", manifest.get("run_id")),
                ("git_sha", manifest.get("git_sha")),
                ("consensus_model", consensus_model),
                ("dry_run", manifest.get("dry_run")),
                ("case_count", manifest.get("case_count")),
                ("source_dir", str(run_dir)),
            ],
        ),
        ("Overall", metric_summary_rows("overall", overall)),
    ]
    for tag in present_tags:
        sections.append(
            (
                f"Tag: {tag}",
                metric_summary_rows(tag, tag_aggregate(aggregate, tag)),
            )
        )

    row = 3
    for title, summary_rows in sections:
        ws.cell(row, 1, title).font = Font(bold=True, color="FFFFFF")
        ws.cell(row, 1).fill = trial_fill
        ws.cell(row, 2, "").fill = trial_fill
        row += 1
        for key, value in summary_rows:
            ws.cell(row, 1, key).fill = label_fill
            ws.cell(row, 1).font = label_font
            ws.cell(row, 2, value)
            row += 1
        row += 1


def build_run_spreadsheet(
    run_dir: Path,
    *,
    output_path: Path | None = None,
) -> Path:
    """Build team_review.xlsx for a single tie-break run directory."""
    run_dir = Path(run_dir)
    output_path = Path(output_path) if output_path is not None else run_dir / "team_review.xlsx"

    manifest = json.loads((run_dir / "manifest.json").read_text())
    records = load_jsonl(run_dir / "records.jsonl")
    aggregate = load_aggregate(run_dir / "aggregate.csv")
    consensus_model = (manifest.get("model_tags") or {}).get("consensus", "?")

    wb = Workbook()
    default = wb.active
    wb.remove(default)
    summary = wb.create_sheet("Summary", 0)

    present_tags: list[str] = []
    for tag, sheet_name in SHEET_BY_TAG:
        tagged = filter_records_by_tag(records, tag)
        if not tagged:
            continue
        present_tags.append(tag)
        cases = pair_cases(tagged)
        ws = wb.create_sheet(sheet_name[:31])
        write_cases(ws, cases, consensus_model=consensus_model)

    write_summary(
        ws=summary,
        manifest=manifest,
        run_dir=run_dir,
        aggregate=aggregate,
        present_tags=present_tags,
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--run-dir",
        type=Path,
        required=True,
        help="Single tie-break run directory (manifest/records/aggregate).",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=None,
        help="Output .xlsx path (default: <run-dir>/team_review.xlsx).",
    )
    args = parser.parse_args()
    path = build_run_spreadsheet(args.run_dir, output_path=args.output)
    print(path)


if __name__ == "__main__":
    main()
