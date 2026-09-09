import json
import logging
from pathlib import Path

import pytest

from experiments.paper.runners import run_bias_1_vs_3


def _minimal_bias_config(tmp_path: Path) -> Path:
    fixture_path = tmp_path / "fixture.json"
    fixture_path.write_text(
        json.dumps(
            {
                "items": [
                    {
                        "trial_id": "trial-1",
                        "trial_pool": "biology",
                        "profile_id": "mtb-h37rv",
                        "gene_id": "Rv0001",
                        "gene_name": "dnaA",
                        "pmc_id": "PMC1",
                        "section": "results",
                        "excerpt_text": "short excerpt",
                    }
                ],
            }
        )
    )
    config_path = tmp_path / "config.yaml"
    config_path.write_text(
        "experiment_id: bias-1-vs-3-small\n"
        "n_trials: 1\n"
        "fixtures:\n"
        f"  papers: {fixture_path}\n"
        "models:\n"
        "  extractors: [model-a, model-b, model-c]\n"
        "  consensus: model-d\n"
    )
    return config_path


def test_dry_run_spreadsheet_mocked_builder_keeps_science_artifacts(
    tmp_path,
    monkeypatch,
):
    paper_dir = tmp_path / "paper"
    monkeypatch.setattr(run_bias_1_vs_3, "PAPER_DIR", paper_dir)
    config_path = _minimal_bias_config(tmp_path)
    built: list[Path] = []

    def fake_builder(run_dir: Path) -> Path:
        out = run_dir / "team_review.xlsx"
        out.write_bytes(b"fake-xlsx")
        built.append(run_dir)
        return out

    monkeypatch.setattr(
        "experiments.paper.runners.run_bias_1_vs_3._bias_spreadsheet_builder",
        fake_builder,
    )

    output_dir = run_bias_1_vs_3.run_bias_experiment(
        config_path=config_path,
        run_id="dry-spreadsheet",
        dry_run=True,
        spreadsheet=True,
    )

    assert (output_dir / "manifest.json").is_file()
    assert (output_dir / "records.jsonl").is_file()
    assert (output_dir / "aggregate.csv").is_file()
    assert built == [output_dir]
    assert (output_dir / "team_review.xlsx").is_file()


def test_spreadsheet_runs_after_derives(tmp_path, monkeypatch):
    paper_dir = tmp_path / "paper"
    monkeypatch.setattr(run_bias_1_vs_3, "PAPER_DIR", paper_dir)
    config_path = _minimal_bias_config(tmp_path)

    split_out = paper_dir / "results" / "split-vs-not" / "split_from_dry-ss"
    order: list[str] = []

    def fake_split(*, bias_run_dir, run_id=None, config_path=None):
        order.append("derive")
        split_out.mkdir(parents=True, exist_ok=True)
        (split_out / "aggregate.csv").write_text(
            "scope,field,n_field_values,split_rate,unanimous_rate,partial_rate\n"
            "overall,,1,0.0,1.0,0.0\n"
        )
        return split_out

    def fake_builder(run_dir: Path) -> Path:
        order.append("spreadsheet")
        manifest = json.loads((run_dir / "manifest.json").read_text())
        assert manifest.get("derived", {}).get("split") == str(split_out)
        out = run_dir / "team_review.xlsx"
        out.write_bytes(b"fake-xlsx")
        return out

    monkeypatch.setattr(
        "experiments.paper.runners.derive_split_vs_not.derive_split_vs_not",
        fake_split,
    )
    monkeypatch.setattr(
        "experiments.paper.runners.run_bias_1_vs_3._bias_spreadsheet_builder",
        fake_builder,
    )

    output_dir = run_bias_1_vs_3.run_bias_experiment(
        config_path=config_path,
        run_id="dry-ss",
        dry_run=True,
        derive_split=True,
        spreadsheet=True,
    )

    assert order == ["derive", "spreadsheet"]
    assert (output_dir / "team_review.xlsx").is_file()


def test_spreadsheet_strict_reraises_builder_failure(tmp_path, monkeypatch):
    paper_dir = tmp_path / "paper"
    monkeypatch.setattr(run_bias_1_vs_3, "PAPER_DIR", paper_dir)
    config_path = _minimal_bias_config(tmp_path)

    def boom(_run_dir: Path) -> Path:
        raise RuntimeError("builder exploded")

    monkeypatch.setattr(
        "experiments.paper.runners.run_bias_1_vs_3._bias_spreadsheet_builder",
        boom,
    )

    with pytest.raises(RuntimeError, match="builder exploded"):
        run_bias_1_vs_3.run_bias_experiment(
            config_path=config_path,
            run_id="dry-spreadsheet-strict",
            dry_run=True,
            spreadsheet=True,
            spreadsheet_strict=True,
        )


def test_spreadsheet_warns_and_continues(tmp_path, monkeypatch, caplog):
    paper_dir = tmp_path / "paper"
    monkeypatch.setattr(run_bias_1_vs_3, "PAPER_DIR", paper_dir)
    config_path = _minimal_bias_config(tmp_path)

    def boom(_run_dir: Path) -> Path:
        raise RuntimeError("builder exploded")

    monkeypatch.setattr(
        "experiments.paper.runners.run_bias_1_vs_3._bias_spreadsheet_builder",
        boom,
    )

    with caplog.at_level(logging.WARNING):
        output_dir = run_bias_1_vs_3.run_bias_experiment(
            config_path=config_path,
            run_id="dry-spreadsheet-warn",
            dry_run=True,
            spreadsheet=True,
            spreadsheet_strict=False,
        )

    assert (output_dir / "aggregate.csv").is_file()
    assert not (output_dir / "team_review.xlsx").exists()
    assert any(
        record.levelno == logging.WARNING and "spreadsheet" in record.message.lower()
        for record in caplog.records
    )


def test_parse_args_accepts_spreadsheet_flags(monkeypatch):
    import sys

    monkeypatch.setattr(
        sys,
        "argv",
        [
            "prog",
            "--config",
            "x.yaml",
            "--spreadsheet",
            "--spreadsheet-strict",
            "--dry-run",
        ],
    )
    args = run_bias_1_vs_3._parse_args()
    assert args.spreadsheet is True
    assert args.spreadsheet_strict is True
    assert args.dry_run is True


def test_build_run_spreadsheet_uses_derived_aggregates(tmp_path):
    from openpyxl import load_workbook

    from experiments.paper.scripts.build_chunking_team_review_spreadsheet import (
        build_run_spreadsheet,
    )

    run_dir = tmp_path / "bias-run"
    run_dir.mkdir()
    split_dir = tmp_path / "split"
    cost_dir = tmp_path / "cost"
    split_dir.mkdir()
    cost_dir.mkdir()
    (split_dir / "aggregate.csv").write_text(
        "scope,field,n_field_values,split_rate,unanimous_rate,partial_rate\n"
        "overall,,2,0.5,0.5,0.0\n"
    )
    (cost_dir / "aggregate.csv").write_text(
        "condition,mean_wall_time_sec\n"
        "consensus_D,1.5\n"
    )
    (run_dir / "manifest.json").write_text(
        json.dumps(
            {
                "experiment_id": "bias-1-vs-3-small",
                "run_id": "chunk-ss",
                "model_tags": {"extractors": ["a", "b", "c"], "consensus": "d"},
                "n_fixture_trials": 1,
                "n_run_trials": 1,
                "derived": {
                    "split": str(split_dir),
                    "cost": str(cost_dir),
                },
            }
        )
    )
    (run_dir / "records.jsonl").write_text(
        json.dumps(
            {
                "record_type": "trial_observable",
                "trial_id": "t1",
                "fixture_trial_id": "t1",
                "trial_pool": "biology",
                "profile_id": "mtb-h37rv",
                "gene_id": "Rv0001",
                "gene_name": "dnaA",
                "pmc_id": "PMC1",
                "section": "results",
                "excerpt_text": "abc",
                "excerpt_preparation": {
                    "tier": "pass",
                    "part_index": 1,
                    "part_count": 1,
                    "chars": 3,
                },
                "outputs": {
                    "extractor_A": {"function": "x"},
                    "extractor_B": None,
                    "extractor_C": None,
                    "consensus_D": {"function": "x"},
                    "single_A": None,
                    "single_B": None,
                    "single_C": None,
                },
                "condition_metrics": {
                    "extractor_A": {"model": "a"},
                    "consensus_D": {"model": "d"},
                },
                "prompts": {},
            }
        )
        + "\n"
    )

    out = build_run_spreadsheet(run_dir)
    assert out == run_dir / "team_review.xlsx"
    assert out.is_file()

    wb = load_workbook(out)
    assert "Summary" in wb.sheetnames
    assert any(name.startswith("Chunk |") for name in wb.sheetnames)
    summary = wb["Summary"]
    values = [
        summary.cell(row, 1).value
        for row in range(1, summary.max_row + 1)
    ]
    assert "Split aggregate" in values
    assert "Cost aggregate" in values
    assert any(
        summary.cell(row, 2).value == str(split_dir / "aggregate.csv")
        for row in range(1, summary.max_row + 1)
        if summary.cell(row, 1).value == "split_aggregate_path"
    )


def test_build_run_spreadsheet_non_chunk_sheet_titles(tmp_path):
    from openpyxl import load_workbook

    from experiments.paper.scripts.build_chunking_team_review_spreadsheet import (
        build_run_spreadsheet,
    )

    run_dir = tmp_path / "bias-run-plain"
    run_dir.mkdir()
    (run_dir / "manifest.json").write_text(
        json.dumps(
            {
                "experiment_id": "bias-1-vs-3-small",
                "run_id": "plain",
                "model_tags": {"extractors": ["a"], "consensus": "d"},
                "n_fixture_trials": 1,
                "n_run_trials": 1,
            }
        )
    )
    (run_dir / "records.jsonl").write_text(
        json.dumps(
            {
                "record_type": "trial_observable",
                "trial_id": "t1",
                "fixture_trial_id": "t1",
                "trial_pool": "biology",
                "profile_id": "ecoli-k12-mg1655",
                "gene_id": "b0001",
                "gene_name": "thrL",
                "pmc_id": "PMC2",
                "section": "results",
                "excerpt_text": "plain excerpt",
                "outputs": {
                    "extractor_A": {"function": "y"},
                    "extractor_B": None,
                    "extractor_C": None,
                    "consensus_D": {"function": "y"},
                    "single_A": None,
                    "single_B": None,
                    "single_C": None,
                },
                "condition_metrics": {
                    "extractor_A": {"model": "a"},
                    "consensus_D": {"model": "d"},
                },
                "prompts": {},
            }
        )
        + "\n"
    )

    out = build_run_spreadsheet(run_dir)
    wb = load_workbook(out)
    assert "Summary" in wb.sheetnames
    assert "E. coli" in wb.sheetnames
    assert not any(name.startswith("Chunk |") for name in wb.sheetnames)
    assert wb["Summary"].cell(1, 1).value == "Bias team review — Summary"
