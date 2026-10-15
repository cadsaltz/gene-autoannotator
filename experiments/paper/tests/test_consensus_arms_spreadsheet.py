import csv
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from experiments.paper.scripts import build_consensus_arms_spreadsheet as builder
from experiments.paper.scripts import compare_consensus_arms as compare
from experiments.paper.scripts import consensus_arms_review as review

ALLOWED_SYMBOLS = set("✔◐✖○✓")
TRIAL_ID = "t001_b0001_PMC1_results#1"
EXCERPT = (
    "Intro line.\n  Gene b0001 is essential   for growth in vitro.\n"
    "It also binds ATP … and more text."
)
TRIO = ("gemma4:12b", "phi4:14b", "dolphin3:8b")


def _label_doc() -> dict:
    return {
        "trial_id": TRIAL_ID,
        "excerpt_note": "synthetic",
        "fields": {
            "function": {
                "gold_status": "present",
                "gold_value": "binds ATP",
                "gold_quote": "binds ATP ... and more",
                "labels": {
                    "gemma4:12b": {"label": "correct"},
                    "phi4:14b": {"label": "partial", "note": "padded"},
                },
            },
            "essential_in_vitro": {
                "gold_status": "present",
                "gold_value": True,
                "gold_quote": "is essential for growth in vitro",
                "labels": {},
            },
            "essential_in_vivo": {
                "gold_status": "absent",
                "gold_value": None,
                "gold_quote": None,
                "labels": {"dolphin3:8b": {"label": "unsupported", "note": "made up"}},
            },
            "infection_impact": {
                "gold_status": "absent",
                "gold_value": None,
                "gold_quote": None,
                "labels": {},
            },
        },
    }


def _record(arm_output: dict, trial_id: str = TRIAL_ID) -> dict:
    return {
        "trial_id": trial_id,
        "excerpt_text": EXCERPT,
        "candidates": {
            "gemma4:12b": {"function": "binds ATP", "essential_in_vitro": None},
            "phi4:14b": {"function": "binds ATP and is padded", "essential_in_vitro": None},
            "dolphin3:8b": {"function": None, "essential_in_vivo": True},
        },
        "arm_output": arm_output,
        "decision_metrics": {"llm_duration_sec": 0.0},
    }


def _write_arm(run_dir: Path, arm: str, record: dict | list[dict]) -> Path:
    arm_dir = run_dir / f"arm_{arm.replace(':', '_')}"
    arm_dir.mkdir(parents=True)
    records = record if isinstance(record, list) else [record]
    (arm_dir / "records.jsonl").write_text("\n".join(json.dumps(r) for r in records) + "\n")
    (arm_dir / "manifest.json").write_text(
        json.dumps({"arm": arm, "models": list(TRIO), "status": "complete"})
    )
    return arm_dir


@pytest.fixture
def labels_dir(tmp_path: Path) -> Path:
    d = tmp_path / "labels"
    d.mkdir()
    (d / "t001_b0001_PMC1_results_1.json").write_text(json.dumps(_label_doc()))
    return d


# --- character preservation (TDD target) ---------------------------------


def test_styled_excerpt_preserves_every_character():
    quotes = ["binds ATP ... and more", "is essential for growth in vitro"]
    rt, found = review.styled_excerpt(EXCERPT, quotes)
    assert rt is not None
    assert found == 3
    assert str(rt) == EXCERPT


def test_styled_excerpt_no_needles_returns_none():
    rt, found = review.styled_excerpt(EXCERPT, [None, ""])
    assert rt is None and found == 0


def test_styled_excerpt_unmatched_quote_changes_nothing():
    rt, found = review.styled_excerpt(EXCERPT, ["text that is not present"])
    assert rt is None and found == 0


# --- flag semantics ------------------------------------------------------


def test_evaluate_arm_output_classes():
    rec = _record({"function": "binds ATP", "essential_in_vivo": True, "essential_in_vitro": None})
    res = review.evaluate_arm_output(_label_doc(), rec["candidates"], rec["arm_output"])
    assert res["classes"]["function"] == "REAL"
    assert res["classes"]["essential_in_vivo"] == "FP"
    assert res["missed"] == ["essential_in_vitro"]
    assert res["quiet"] == ["infection_impact"]


def test_evaluate_arm_output_unlabeled_value_is_not_guessed():
    rec = _record({"function": "a brand new merged sentence"})
    res = review.evaluate_arm_output(_label_doc(), rec["candidates"], rec["arm_output"])
    assert res["classes"]["function"] == "UNLABELED"


def test_flag_text_uses_only_allowed_symbols():
    rec = _record({"function": "binds ATP", "essential_in_vivo": True})
    res = review.evaluate_arm_output(_label_doc(), rec["candidates"], rec["arm_output"])
    text = review.flag_text(res, has_needles=True)
    assert "✔" in text and "✖" in text and "○" in text
    assert {ch for ch in text if ord(ch) > 0x2000} <= ALLOWED_SYMBOLS


# --- workbook ------------------------------------------------------------


def test_builder_creates_required_rows_and_preserves_excerpt(tmp_path, labels_dir):
    arm_dir = _write_arm(
        tmp_path / "run",
        "A",
        _record({"function": "binds ATP", "essential_in_vivo": True, "essential_in_vitro": None}),
    )
    out = tmp_path / "out" / "arm_A.xlsx"
    builder.build_workbook(arm_dir, labels_dir, out)

    wb = load_workbook(out, rich_text=True)
    assert wb.sheetnames[0] == "README"
    ws = wb["Arm A"]
    labels = {str(ws.cell(r, 1).value): r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value}
    for required in (
        "TRIAL",
        "trial_id",
        "EXCERPT",
        "CANDIDATE gemma4:12b",
        "CANDIDATE phi4:14b",
        "CANDIDATE dolphin3:8b",
        "ARM OUTPUT (A)",
    ):
        assert required in labels, required

    excerpt_cell = ws.cell(labels["EXCERPT"], 2).value
    assert str(excerpt_cell) == EXCERPT
    assert ws.cell(labels["trial_id"], 2).value == TRIAL_ID

    flags = str(ws.cell(labels["ARM OUTPUT (A)"], 7).value)
    assert "✔" in flags and "✖" in flags and "○" in flags
    assert {ch for ch in flags if ord(ch) > 0x2000} <= ALLOWED_SYMBOLS
    cand_flags = str(ws.cell(labels["CANDIDATE phi4:14b"], 7).value)
    assert "◐" in cand_flags


def test_builder_truncates_only_past_excel_cell_limit(tmp_path, labels_dir):
    long_text = "x" * 40000
    rec = _record({"function": None}, trial_id="t999_unlabeled#1")
    rec["excerpt_text"] = long_text
    arm_dir = _write_arm(tmp_path / "run", "A", rec)
    out = tmp_path / "long.xlsx"
    builder.build_workbook(arm_dir, labels_dir, out)
    ws = load_workbook(out, rich_text=True)["Arm A"]
    row = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == "EXCERPT")
    value = str(ws.cell(row, 2).value)
    assert len(value) <= 32767 and long_text.startswith(value)
    assert "no label file" in str(ws.cell(row, 7).value).lower()


# --- compare -------------------------------------------------------------


def test_compare_writes_scoreboard_and_vs_a(tmp_path, labels_dir):
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    _write_arm(run_dir, "A", _record({"function": "binds ATP", "essential_in_vivo": True}))
    _write_arm(
        run_dir,
        "B",
        _record({"function": None, "essential_in_vitro": None, "essential_in_vivo": None}),
    )
    out = tmp_path / "compare"
    assert compare.main(
        ["--run-dir", str(run_dir), "--labels", str(labels_dir), "--out", str(out)]
    ) == 0

    rows = list(csv.DictReader((out / "scoreboard.csv").open()))
    by_arm = {r["arm"]: r for r in rows}
    assert by_arm["A"]["needles_total"] == "2"
    assert by_arm["A"]["found"] == "1" and by_arm["A"]["missed"] == "1"
    assert by_arm["A"]["fp"] == "1"
    assert by_arm["B"]["found"] == "0" and by_arm["B"]["missed"] == "2"
    assert by_arm["B"]["fp"] == "0"

    md = (out / "vs_A.md").read_text()
    assert "## Arm B vs A" in md
    assert "function" in md and "lost" in md.lower()
    assert "essential_in_vivo" in md  # FP removed vs A


def test_compare_requires_arm_a(tmp_path, labels_dir):
    run_dir = tmp_path / "run"
    run_dir.mkdir()
    _write_arm(run_dir, "B", _record({"function": None}))
    with pytest.raises(SystemExit):
        compare.main(
            ["--run-dir", str(run_dir), "--labels", str(labels_dir), "--out", str(tmp_path / "o")]
        )
