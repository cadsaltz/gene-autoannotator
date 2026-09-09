import csv
import json
from pathlib import Path

import pytest

from experiments.paper.runners.run_tiebreak_consensus import (
    NONSENSE_EXPERIMENT_TAG,
    _aggregate_row,
    _aggregate_rows,
    run_experiment,
)


CONFIG = Path("experiments/paper/configs/tiebreak-non-nonsense.yaml")
COMBINED_CONFIG = Path("experiments/paper/configs/tiebreak-consensus.yaml")
NON_NONSENSE_EXPERIMENT_TAG = "tiebreak-non-nonsense"


def test_dry_run_writes_scored_artifacts_without_ollama(tmp_path, monkeypatch):
    def fail_if_called(*args, **kwargs):
        raise AssertionError("dry-run must not call Ollama")

    monkeypatch.setattr(
        "experiments.paper.runners.general_consensus.ollama_chat",
        fail_if_called,
    )

    run_dir = run_experiment(
        config_path=CONFIG,
        run_id="drytest",
        dry_run=True,
        limit=5,
        results_root=tmp_path,
    )

    assert (run_dir / "manifest.json").is_file()
    assert (run_dir / "records.jsonl").is_file()
    assert (run_dir / "aggregate.csv").is_file()

    manifest = json.loads((run_dir / "manifest.json").read_text())
    assert manifest["model_tags"]["consensus"] == "qwen3.5:27b"
    assert manifest["dry_run"] is True
    assert manifest["case_selection"] == {
        "include_tags": ["tiebreak-non-nonsense"],
        "match_any": True,
    }
    assert manifest["case_counts_by_tag"] == {"tiebreak-non-nonsense": 5}

    records = [
        json.loads(line)
        for line in (run_dir / "records.jsonl").read_text().splitlines()
    ]
    assert len(records) == 10
    assert {record["condition"] for record in records} == {
        "no_consensus_pick_extractor_0",
        "hybrid_consensus",
    }
    assert all(
        {"expected", "observed", "llm_invoked", "provenance"} <= record.keys()
        for record in records
    )
    assert any(
        record["condition"] == "hybrid_consensus" and record["llm_invoked"]
        for record in records
    )

    aggregates = list(csv.DictReader((run_dir / "aggregate.csv").open()))
    assert aggregates
    assert {
        "consensus_match_soft_rate",
        "baseline_match_soft_rate",
        "necessity_delta",
        "llm_invoked_rate",
        "expect_llm_calibration_rate",
        "invention_rate",
    } <= aggregates[0].keys()


def _make_record(
    *,
    case_id: str,
    condition: str,
    experiment_tags: list[str],
    case_family: str = "exact_majority",
    match_soft: bool = True,
) -> dict:
    return {
        "case_id": case_id,
        "condition": condition,
        "experiment_tags": experiment_tags,
        "case_family": case_family,
        "domain": "general",
        "expect_llm": False,
        "match_exact": match_soft,
        "match_soft": match_soft,
        "llm_invoked": False,
        "expect_llm_matched": True,
        "invention": False,
    }


def test_aggregate_row_nonsense_metric_uses_tag_not_family():
    """Nonsense adoption rate must not be diluted by non-nonsense tag rows."""
    records = [
        _make_record(
            case_id="nn-1",
            condition="no_consensus_pick_extractor_0",
            experiment_tags=["tiebreak-non-nonsense"],
            match_soft=False,
        ),
        _make_record(
            case_id="nn-1",
            condition="hybrid_consensus",
            experiment_tags=["tiebreak-non-nonsense"],
            match_soft=True,
        ),
        _make_record(
            case_id="ns-1",
            condition="no_consensus_pick_extractor_0",
            experiment_tags=[NONSENSE_EXPERIMENT_TAG],
            case_family="exact_majority_nonsense",
            match_soft=False,
        ),
        _make_record(
            case_id="ns-1",
            condition="hybrid_consensus",
            experiment_tags=[NONSENSE_EXPERIMENT_TAG],
            case_family="exact_majority_nonsense",
            match_soft=True,
        ),
    ]
    overall = _aggregate_row(records, scope="overall", value="all")
    assert overall["case_count"] == 2
    assert overall["nonsense_majority_adoption_rate"] == 1.0

    non_nonsense = _aggregate_row(
        [record for record in records if "tiebreak-non-nonsense" in record["experiment_tags"]],
        scope="experiment_tag",
        value="tiebreak-non-nonsense",
    )
    assert non_nonsense["nonsense_majority_adoption_rate"] == 0.0

    nonsense = _aggregate_row(
        [record for record in records if NONSENSE_EXPERIMENT_TAG in record["experiment_tags"]],
        scope="experiment_tag",
        value=NONSENSE_EXPERIMENT_TAG,
    )
    assert nonsense["nonsense_majority_adoption_rate"] == 1.0


def test_aggregate_rows_emits_experiment_tag_scope():
    records = [
        _make_record(
            case_id="a",
            condition="hybrid_consensus",
            experiment_tags=["tiebreak-non-nonsense"],
        ),
        _make_record(
            case_id="b",
            condition="hybrid_consensus",
            experiment_tags=[NONSENSE_EXPERIMENT_TAG],
            case_family="exact_majority_nonsense",
        ),
    ]
    rows = _aggregate_rows(records)
    tag_rows = [row for row in rows if row["scope"] == "experiment_tag"]
    assert {row["value"] for row in tag_rows} == {
        "tiebreak-non-nonsense",
        NONSENSE_EXPERIMENT_TAG,
    }


def test_dry_run_combined_config_emits_tag_scoped_aggregates(tmp_path, monkeypatch):
    def fail_if_called(*args, **kwargs):
        raise AssertionError("dry-run must not call Ollama")

    monkeypatch.setattr(
        "experiments.paper.runners.general_consensus.ollama_chat",
        fail_if_called,
    )

    run_dir = run_experiment(
        config_path=COMBINED_CONFIG,
        run_id="combined-drytest",
        dry_run=True,
        limit=10,
        results_root=tmp_path,
    )

    manifest = json.loads((run_dir / "manifest.json").read_text())
    assert manifest["case_selection"] == {
        "include_tags": ["tiebreak-non-nonsense", "tiebreak-nonsense"],
        "match_any": True,
    }
    assert set(manifest["case_counts_by_tag"]) == {
        "tiebreak-non-nonsense",
        "tiebreak-nonsense",
    }

    records = [
        json.loads(line)
        for line in (run_dir / "records.jsonl").read_text().splitlines()
    ]
    record_tags = {
        tag
        for record in records
        for tag in record.get("experiment_tags", [])
    }
    assert "tiebreak-non-nonsense" in record_tags
    assert "tiebreak-nonsense" in record_tags

    aggregates = list(csv.DictReader((run_dir / "aggregate.csv").open()))
    tag_rows = [row for row in aggregates if row["scope"] == "experiment_tag"]
    assert {row["value"] for row in tag_rows} == {
        "tiebreak-non-nonsense",
        "tiebreak-nonsense",
    }
    nonsense_row = next(
        row for row in tag_rows if row["value"] == NONSENSE_EXPERIMENT_TAG
    )
    non_nonsense_row = next(
        row for row in tag_rows if row["value"] == "tiebreak-non-nonsense"
    )
    assert float(nonsense_row["nonsense_majority_adoption_rate"]) > 0.0
    assert float(non_nonsense_row["nonsense_majority_adoption_rate"]) == 0.0


def test_dry_run_spreadsheet_mocked_builder_keeps_science_artifacts(tmp_path, monkeypatch):
    built: list[Path] = []

    def fake_builder(run_dir: Path) -> Path:
        out = run_dir / "team_review.xlsx"
        out.write_bytes(b"fake-xlsx")
        built.append(run_dir)
        return out

    def fail_if_called(*args, **kwargs):
        raise AssertionError("dry-run must not call Ollama")

    monkeypatch.setattr(
        "experiments.paper.runners.general_consensus.ollama_chat",
        fail_if_called,
    )
    monkeypatch.setattr(
        "experiments.paper.runners.run_tiebreak_consensus._tiebreak_spreadsheet_builder",
        fake_builder,
    )

    run_dir = run_experiment(
        config_path=CONFIG,
        run_id="dry-spreadsheet",
        dry_run=True,
        limit=3,
        results_root=tmp_path,
        spreadsheet=True,
    )

    assert (run_dir / "manifest.json").is_file()
    assert (run_dir / "records.jsonl").is_file()
    assert (run_dir / "aggregate.csv").is_file()
    assert built == [run_dir]
    assert (run_dir / "team_review.xlsx").is_file()


def test_spreadsheet_strict_reraises_builder_failure(tmp_path, monkeypatch):
    def fail_if_called(*args, **kwargs):
        raise AssertionError("dry-run must not call Ollama")

    monkeypatch.setattr(
        "experiments.paper.runners.general_consensus.ollama_chat",
        fail_if_called,
    )

    def boom(_run_dir: Path) -> Path:
        raise RuntimeError("builder exploded")

    monkeypatch.setattr(
        "experiments.paper.runners.run_tiebreak_consensus._tiebreak_spreadsheet_builder",
        boom,
    )

    with pytest.raises(RuntimeError, match="builder exploded"):
        run_experiment(
            config_path=CONFIG,
            run_id="dry-spreadsheet-strict",
            dry_run=True,
            limit=2,
            results_root=tmp_path,
            spreadsheet=True,
            spreadsheet_strict=True,
        )


def _write_minimal_run(run_dir: Path, *, tags_by_case: dict[str, list[str]]) -> None:
    run_dir.mkdir(parents=True, exist_ok=True)
    records = []
    for case_id, tags in tags_by_case.items():
        for condition in ("no_consensus_pick_extractor_0", "hybrid_consensus"):
            records.append(
                {
                    "case_id": case_id,
                    "domain": "general",
                    "case_family": "exact_majority",
                    "experiment_tags": tags,
                    "field_key": "label",
                    "condition": condition,
                    "candidates": [{"label": "a"}, {"label": "a"}, {"label": "b"}],
                    "expected": "a",
                    "observed": "a",
                    "llm_invoked": False,
                    "provenance": "extractor_0",
                    "match_exact": True,
                    "match_soft": True,
                    "invention": False,
                    "expect_llm": False,
                    "expect_llm_matched": True,
                }
            )
    (run_dir / "records.jsonl").write_text(
        "\n".join(json.dumps(row) for row in records) + "\n"
    )
    (run_dir / "manifest.json").write_text(
        json.dumps(
            {
                "experiment_id": "tiebreak-consensus",
                "run_id": "fixture-run",
                "dry_run": True,
                "git_sha": "deadbeef",
                "case_count": len(tags_by_case),
                "model_tags": {"consensus": "qwen3.5:27b"},
            }
        )
    )
    (run_dir / "aggregate.csv").write_text(
        "scope,value,case_count,consensus_match_soft_rate,"
        "consensus_match_exact_rate,baseline_match_soft_rate,"
        "necessity_delta,llm_invoked_rate,expect_llm_calibration_rate,"
        "invention_rate,nonsense_majority_adoption_rate\n"
        "overall,all,2,1.0,1.0,1.0,0.0,0.0,1.0,0.0,0.0\n"
        "experiment_tag,tiebreak-non-nonsense,1,1.0,1.0,1.0,0.0,0.0,1.0,0.0,0.0\n"
        "experiment_tag,tiebreak-nonsense,1,1.0,1.0,1.0,0.0,0.0,1.0,0.0,1.0\n"
    )


def test_build_run_spreadsheet_combined_sheets(tmp_path):
    from openpyxl import load_workbook

    from experiments.paper.scripts.build_tiebreak_team_review_spreadsheet import (
        build_run_spreadsheet,
    )

    run_dir = tmp_path / "combined"
    _write_minimal_run(
        run_dir,
        tags_by_case={
            "nn-1": [NON_NONSENSE_EXPERIMENT_TAG],
            "ns-1": [NONSENSE_EXPERIMENT_TAG],
        },
    )
    out = build_run_spreadsheet(run_dir)
    assert out == run_dir / "team_review.xlsx"
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "Non-nonsense", "Nonsense"]
    summary = wb["Summary"]
    summary_text = " ".join(
        str(cell.value)
        for row in summary.iter_rows(max_row=40, max_col=2)
        for cell in row
        if cell.value is not None
    )
    assert "overall" in summary_text.lower() or "Overall" in summary_text
    assert NON_NONSENSE_EXPERIMENT_TAG in summary_text
    assert NONSENSE_EXPERIMENT_TAG in summary_text
    nn = wb["Non-nonsense"]
    assert any(
        cell.value and "CANDIDATES" in str(cell.value)
        for row in nn.iter_rows(max_col=1)
        for cell in row
    )
    assert any(
        cell.value and "CONSENSUS PROMPT" in str(cell.value)
        for row in nn.iter_rows(max_col=1)
        for cell in row
    )
    assert any(
        cell.value and "baseline" in str(cell.value).lower()
        for row in nn.iter_rows(max_col=1)
        for cell in row
    )


def test_build_run_spreadsheet_omits_absent_tag_sheet(tmp_path):
    from openpyxl import load_workbook

    from experiments.paper.scripts.build_tiebreak_team_review_spreadsheet import (
        build_run_spreadsheet,
    )

    run_dir = tmp_path / "nn-only"
    _write_minimal_run(
        run_dir,
        tags_by_case={"nn-1": [NON_NONSENSE_EXPERIMENT_TAG]},
    )
    # Override aggregate to non-nonsense-only.
    (run_dir / "aggregate.csv").write_text(
        "scope,value,case_count,consensus_match_soft_rate,"
        "consensus_match_exact_rate,baseline_match_soft_rate,"
        "necessity_delta,llm_invoked_rate,expect_llm_calibration_rate,"
        "invention_rate,nonsense_majority_adoption_rate\n"
        "overall,all,1,1.0,1.0,1.0,0.0,0.0,1.0,0.0,0.0\n"
        "experiment_tag,tiebreak-non-nonsense,1,1.0,1.0,1.0,0.0,0.0,1.0,0.0,0.0\n"
    )
    out = build_run_spreadsheet(run_dir)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "Non-nonsense"]
